from io import BytesIO
from pathlib import Path

from fastapi import HTTPException, UploadFile
from PIL import Image
from PyPDF2 import PdfReader
from docx import Document
from pptx import Presentation


ALLOWED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".pptx",
    ".png",
    ".jpg",
    ".jpeg",
    ".webp",
}

IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp"}


def validate_upload(upload: UploadFile, max_bytes: int) -> str:
    filename = upload.filename or ""
    if not filename.strip():
        raise HTTPException(status_code=400, detail="Uploaded file is missing a filename.")

    suffix = Path(filename).suffix.lower()
    if suffix not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=415,
            detail="Unsupported file type. Use PDF, DOCX, PPTX, PNG, JPG/JPEG, or WEBP.",
        )

    if upload.size is not None and upload.size > max_bytes:
        raise HTTPException(
            status_code=413,
            detail=f"File is too large. Maximum allowed size is {max_bytes // (1024 * 1024)} MB.",
        )

    return suffix


async def read_upload(upload: UploadFile, max_bytes: int) -> bytes:
    data = await upload.read()
    if not data:
        raise HTTPException(status_code=400, detail="Uploaded file is empty.")
    if len(data) > max_bytes:
        raise HTTPException(
            status_code=413,
            detail=f"File is too large. Maximum allowed size is {max_bytes // (1024 * 1024)} MB.",
        )
    return data


def _finalize_extracted_text(text: str) -> str:
    cleaned = text.strip()
    if not cleaned:
        raise HTTPException(
            status_code=422,
            detail="No readable study content was found in the uploaded file.",
        )
    return cleaned


def extract_text(data: bytes, suffix: str) -> str:
    try:
        if suffix == ".pdf":
            reader = PdfReader(BytesIO(data))
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
            return _finalize_extracted_text(text)

        if suffix == ".docx":
            document = Document(BytesIO(data))
            text = "\n".join(p.text for p in document.paragraphs if p.text.strip())
            return _finalize_extracted_text(text)

        if suffix == ".pptx":
            presentation = Presentation(BytesIO(data))
            chunks: list[str] = []
            for slide in presentation.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        chunks.append(shape.text.strip())
            return _finalize_extracted_text("\n".join(chunks))

        if suffix in IMAGE_EXTENSIONS:
            # Validate that the bytes are a real image.
            with Image.open(BytesIO(data)) as image:
                image.verify()
            raise HTTPException(
                status_code=422,
                detail="No readable study content was found in the uploaded image.",
            )

    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(
            status_code=422,
            detail=f"Could not read the uploaded file: {exc}",
        ) from exc

    raise HTTPException(status_code=415, detail="Unsupported file format.")
